"""MC-03 data loading, validation and report orchestration."""

from datetime import date
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from .config import (
    ALL_PROVINCES_TOKEN,
    CODE,
    CODIGOS_ESTANDAR,
    CUTOFF_CELL,
    DEFAULT_PROVINCE,
    EXCEL_SHEET,
    HEADER_ROW,
    MONTHS_2026,
    NAME,
    REGLAS_NEGOCIO,
    REQUIRED_COLUMNS,
    VERIFICATION_MONTHS,
)
from .cred import cred_detail, valid_cred
from .denominator import is_in_denominator
from .screening import tamizaje_detail, valid_tamizaje
from .utils import (
    clean_identifier,
    clean_text,
    coverage_semaphore,
    normalize_filter,
    normalize_target,
    parse_date,
    parse_month_key,
    safe_row_value,
    to_int,
    to_number,
)
from .vaccines import vaccine_detail, valid_vaccine


def load_sample_data(filepath: Path) -> dict[str, Any] | None:
    if not filepath.exists():
        return None

    workbook = load_workbook(filepath, data_only=True, read_only=True)
    if EXCEL_SHEET not in workbook.sheetnames:
        raise ValueError("La hoja 'Detalle_Ate' no se encuentra en el archivo de datos")

    sheet = workbook[EXCEL_SHEET]
    cutoff_date = parse_date(sheet[CUTOFF_CELL].value)

    df = pd.read_excel(
        filepath,
        sheet_name=EXCEL_SHEET,
        header=HEADER_ROW - 1,
        engine="openpyxl",
    )

    return {"data": df, "cutoff_date": cutoff_date}


def validate_data_file(filepath: Path) -> dict[str, Any]:
    if not filepath.exists():
        return {
            "valid": False,
            "errors": ["No se encontro el archivo cargado."],
            "warnings": [],
            "summary": {},
        }

    errors = []
    warnings = []
    summary: dict[str, Any] = {
        "indicator_code": CODE,
        "indicator_name": NAME,
        "filename": filepath.name,
        "file_size_bytes": filepath.stat().st_size,
        "sheet_name": EXCEL_SHEET,
        "cutoff_cell": CUTOFF_CELL,
        "header_row": HEADER_ROW,
        "validation_label": "Obs_Eval",
        "required_columns_count": len(REQUIRED_COLUMNS),
    }

    try:
        workbook = load_workbook(filepath, data_only=True, read_only=True)
    except Exception as exc:
        return {
            "valid": False,
            "errors": [f"No se pudo abrir el Excel: {exc}"],
            "warnings": [],
            "summary": summary,
        }

    if EXCEL_SHEET not in workbook.sheetnames:
        errors.append("No se encontro la hoja obligatoria 'Detalle_Ate'.")
        return {"valid": False, "errors": errors, "warnings": warnings, "summary": summary}

    sheet = workbook[EXCEL_SHEET]
    cutoff_date = parse_date(sheet[CUTOFF_CELL].value)
    if cutoff_date is None:
        warnings.append(f"No se pudo leer una fecha de corte valida desde la celda {CUTOFF_CELL}.")
    summary["cutoff_date"] = cutoff_date

    try:
        df = pd.read_excel(filepath, sheet_name=EXCEL_SHEET, header=HEADER_ROW - 1, engine="openpyxl")
    except Exception as exc:
        return {
            "valid": False,
            "errors": [f"No se pudo leer la hoja 'Detalle_Ate': {exc}"],
            "warnings": warnings,
            "summary": summary,
        }

    missing_columns = [column for column in REQUIRED_COLUMNS if column not in df.columns]
    if missing_columns:
        errors.append("Faltan columnas obligatorias: " + ", ".join(missing_columns))
    summary["missing_columns"] = missing_columns

    total_rows = len(df)
    summary["total_rows"] = total_rows
    summary["total_columns"] = len(df.columns)
    summary["columns_found"] = list(df.columns)

    if total_rows == 0:
        errors.append("La hoja 'Detalle_Ate' no contiene registros.")

    if "Desc_prov" in df.columns:
        provinces = sorted(
            value
            for value in df["Desc_prov"].dropna().astype(str).str.strip().unique().tolist()
            if value
        )
        summary["provinces"] = provinces
    else:
        summary["provinces"] = []

    if "Mes_eva" in df.columns:
        months = sorted(
            value
            for value in df["Mes_eva"].dropna().astype(str).str.strip().unique().tolist()
            if value
        )
        summary["months"] = months
    else:
        summary["months"] = []

    if "Obs_Eval" in df.columns:
        obs_counts = df["Obs_Eval"].fillna("VACIO").astype(str).str.strip().value_counts().to_dict()
        summary["obs_eval_counts"] = obs_counts
        summary["status_counts"] = obs_counts
        normalized_counts = {key.upper(): value for key, value in obs_counts.items()}
        if "EVALUADO" not in normalized_counts:
            warnings.append("No se encontraron registros con Obs_Eval = Evaluado.")
    else:
        summary["obs_eval_counts"] = {}
        summary["status_counts"] = {}

    if "Esta_pac" in df.columns:
        summary["insurance_counts"] = df["Esta_pac"].fillna("VACIO").astype(str).str.strip().value_counts().to_dict()
    else:
        summary["insurance_counts"] = {}

    return {
        "valid": not errors,
        "errors": errors,
        "warnings": warnings,
        "summary": summary,
    }


def filter_data(df: pd.DataFrame, province: str | None = DEFAULT_PROVINCE) -> pd.DataFrame:
    selected_province = normalize_filter(province)
    if selected_province is None or "Desc_prov" not in df.columns:
        return df

    return df[df["Desc_prov"].astype(str).str.strip().str.upper() == selected_province]


def get_filter_options(df: pd.DataFrame) -> dict[str, Any]:
    provinces = []
    if df is not None and not df.empty and "Desc_prov" in df.columns:
        provinces = sorted(
            value
            for value in df["Desc_prov"].dropna().astype(str).str.strip().unique().tolist()
            if value
        )

    return {
        "provinces": provinces,
        "default_province": DEFAULT_PROVINCE,
        "default_target_coverage": REGLAS_NEGOCIO["META_COBERTURA_MENSUAL"],
        "included_insurance_types": sorted(
            "VACIO/SIN SEGURO" if value == "" else value
            for value in REGLAS_NEGOCIO["SEGUROS_INCLUIDOS"]
        ),
        "exclusion_criteria": {
            "included_obs_eval": "Evaluado",
            "excluded_obs_eval": "No_Evaluado",
        },
    }


def _normalize_target(value: float | None) -> float:
    return normalize_target(value, REGLAS_NEGOCIO["META_COBERTURA_MENSUAL"])


def _coverage_semaphore(coverage: float, target: float) -> str:
    return coverage_semaphore(coverage, target)


def _is_in_denominator(row: pd.Series, year: int, month: int) -> bool:
    return is_in_denominator(row, year, month)


def _valid_vaccine(row: pd.Series, date_col: str, result_col: str, age_col: str) -> bool:
    return valid_vaccine(row, date_col, result_col, age_col)


def _valid_cred(row: pd.Series, number: int) -> bool:
    return valid_cred(row, number)


def _valid_tamizaje(row: pd.Series) -> bool:
    return valid_tamizaje(row)


def evaluate_package(row: pd.Series, reference_date: date | None = None) -> dict[str, Any]:
    reference_date = reference_date or date.today()
    details = {
        "bcg": vaccine_detail(row, "BCG", "fec1_BCG", "resul1_BCG", "Edad_ate1_BCG", reference_date),
        "hvb": vaccine_detail(row, "HvB", "fecHVB", "resulHVB", "Edad_ateHVB", reference_date),
        "cred_1": cred_detail(row, 1, reference_date),
        "cred_2": cred_detail(row, 2, reference_date),
        "cred_3": cred_detail(row, 3, reference_date),
        "tamizaje": tamizaje_detail(row, reference_date),
    }
    checks = {key: value["cumple"] for key, value in details.items()}
    reasons = [value["mensaje"] for value in details.values() if not value["cumple"]]
    return {"complete": not reasons, "checks": checks, "details": details, "reasons": reasons}


def _omiso_from_row(row: pd.Series, reasons: list[str]) -> dict[str, Any]:
    year, month, month_name = parse_month_key(safe_row_value(row, "Mes_eva"))
    return {
        "Mes_eva": clean_text(safe_row_value(row, "Mes_eva")) or None,
        "month": month_name,
        "year": year,
        "afi_DNI": clean_text(safe_row_value(row, "afi_DNI")) or None,
        "NumCNV": clean_text(safe_row_value(row, "NumCNV")) or None,
        "fec_Nac": clean_text(safe_row_value(row, "fec_Nac")) or None,
        "afi_nombres": clean_text(safe_row_value(row, "afi_nombres")) or None,
        "afi_appaterno": clean_text(safe_row_value(row, "afi_appaterno")) or None,
        "afi_apmaterno": clean_text(safe_row_value(row, "afi_apmaterno")) or None,
        "Desc_prov": clean_text(safe_row_value(row, "Desc_prov")) or None,
        "Des_MicroRed": clean_text(safe_row_value(row, "Des_MicroRed")) or None,
        "pre_CodigoRENAES": clean_identifier(safe_row_value(row, "pre_CodigoRENAES")) or None,
        "Des_EESS": clean_text(safe_row_value(row, "Des_EESS")) or None,
        "reason": "; ".join(reasons),
    }


def build_report_summary(
    df: pd.DataFrame,
    cutoff_date: date | None = None,
    province: str | None = DEFAULT_PROVINCE,
    target_coverage: float | None = None,
) -> Any:
    df = filter_data(df, province)
    period_start = date(2026, 6, 1)
    period_end = date(2026, 11, 30)
    target = _normalize_target(target_coverage)
    monthly = []
    omisos = []

    for month, month_name in MONTHS_2026:
        year = 2026
        denominator_df = df[df.apply(lambda row: is_in_denominator(row, year, month), axis=1)]

        numerator = 0
        for _, row in denominator_df.iterrows():
            package = evaluate_package(row, cutoff_date)
            if package["complete"]:
                numerator += 1
            else:
                omisos.append(_omiso_from_row(row, package["reasons"]))

        denominator = len(denominator_df)
        coverage = round((numerator / denominator) * 100, 2) if denominator else 0.0
        monthly.append(
            {
                "month": month_name,
                "year": year,
                "in_verification_period": month in VERIFICATION_MONTHS,
                "compliant": denominator > 0 and coverage >= target,
                "semaphore": coverage_semaphore(coverage, target),
                "coverage": coverage,
                "denominator": denominator,
                "numerator": numerator,
            }
        )

    verification_monthly = [item for item in monthly if item["in_verification_period"]]
    months_met = sum(1 for item in verification_monthly if item["compliant"])

    return {
        "period_start": period_start,
        "period_end": period_end,
        "cut_off_date": cutoff_date,
        "target_coverage": target,
        "months_evaluated": len(verification_monthly),
        "months_met": months_met,
        "committed": months_met >= 5,
        "monthly": monthly,
        "omisos": omisos,
    }


def search_by_dni(
    df: pd.DataFrame,
    dni: str,
    reference_date: date | None = None,
    province: str | None = DEFAULT_PROVINCE,
) -> dict[str, Any] | None:
    """Search a newborn record by DNI and return MC-03 compliance data."""
    if df is None or df.empty:
        return None

    df = filter_data(df, province)
    result = df[df["afi_DNI"].astype(str).str.strip() == dni.strip()]
    if result.empty:
        return None

    row = result.iloc[0]
    package = evaluate_package(row, reference_date)

    personal = {
        "afi_DNI": clean_text(row.get("afi_DNI")) or None,
        "NumCNV": clean_text(row.get("NumCNV")) or None,
        "afi_nombres": clean_text(row.get("afi_nombres")) or None,
        "afi_appaterno": clean_text(row.get("afi_appaterno")) or None,
        "afi_apmaterno": clean_text(row.get("afi_apmaterno")) or None,
        "fec_Nac": clean_text(row.get("fec_Nac")) or None,
        "peso": to_number(row.get("peso")),
        "edadGEst": to_int(row.get("edadGEst")),
        "Desc_prov": clean_text(row.get("Desc_prov")) or None,
        "Des_MicroRed": clean_text(row.get("Des_MicroRed")) or None,
        "pre_CodigoRENAES": clean_identifier(row.get("pre_CodigoRENAES")) or None,
        "Des_EESS": clean_text(row.get("Des_EESS")) or None,
    }

    cred_controls = []
    for i in range(1, 4):
        fecha_col = f"Fecha_Atencion_{i}"
        cred_controls.append(
            {
                "numero": i,
                "fecha": clean_text(row.get(fecha_col)) or None,
                "edad_atencion_dias": to_int(row.get(f"Edad_Atencion_{i}")),
                "cumple": package["checks"][f"cred_{i}"],
                "estado": package["details"][f"cred_{i}"]["estado"],
                "mensaje": package["details"][f"cred_{i}"]["mensaje"],
                "fecha_inicio": package["details"][f"cred_{i}"]["fecha_inicio"],
                "fecha_limite": package["details"][f"cred_{i}"]["fecha_limite"],
            }
        )

    return {
        "personal": personal,
        "vacunas": {
            "BCG": {
                "codigo": CODIGOS_ESTANDAR["BCG"],
                "fecha": clean_text(row.get("fec1_BCG")) or None,
                "resultado": clean_text(row.get("resul1_BCG")) or None,
                "edad_atencion_dias": to_int(row.get("Edad_ate1_BCG")),
                "dosis_registradas": 1 if clean_text(row.get("fec1_BCG")) else 0,
                "dosis_evaluadas": 1,
                "cumple": package["checks"]["bcg"],
                "estado": package["details"]["bcg"]["estado"],
                "mensaje": package["details"]["bcg"]["mensaje"],
                "fecha_inicio": package["details"]["bcg"]["fecha_inicio"],
                "fecha_limite": package["details"]["bcg"]["fecha_limite"],
            },
            "HVB": {
                "codigo": CODIGOS_ESTANDAR["HVB"],
                "fecha": clean_text(row.get("fecHVB")) or None,
                "resultado": clean_text(row.get("resulHVB")) or None,
                "edad_atencion_dias": to_int(row.get("Edad_ateHVB")),
                "dosis_registradas": 1 if clean_text(row.get("fecHVB")) else 0,
                "dosis_evaluadas": 1,
                "cumple": package["checks"]["hvb"],
                "estado": package["details"]["hvb"]["estado"],
                "mensaje": package["details"]["hvb"]["mensaje"],
                "fecha_inicio": package["details"]["hvb"]["fecha_inicio"],
                "fecha_limite": package["details"]["hvb"]["fecha_limite"],
            },
        },
        "cred_controls": cred_controls,
        "tamizaje": {
            "fecha": clean_text(row.get("Fecha_Atencion_TN")) or None,
            "edad_atencion_dias": to_int(row.get("Edad_Atencion_TN")),
            "cumple": package["checks"]["tamizaje"],
            "estado": package["details"]["tamizaje"]["estado"],
            "mensaje": package["details"]["tamizaje"]["mensaje"],
            "fecha_inicio": package["details"]["tamizaje"]["fecha_inicio"],
            "fecha_limite": package["details"]["tamizaje"]["fecha_limite"],
        },
        "paquete_completo": package["complete"],
    }
