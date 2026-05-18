import hashlib
import json
import re
import shutil
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path
from uuid import UUID
from uuid import uuid4

import pandas as pd
from fastapi import BackgroundTasks, Depends, FastAPI, File, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from .schemas import (
        DataActivateRequest,
        DataCurrentResponse,
        DataUploadPreviewResponse,
        FilterOptionsResponse,
        HealthResponse,
        LoginRequest,
        LoginResponse,
        AuditEventsResponse,
        OmisosResponse,
        ReportSummary,
        SearchDNIResult,
        UploadHistoryResponse,
    )
    from .services import build_report_summary, get_filter_options, load_sample_data, search_by_dni, validate_data_file
    from .indicators.registry import get_indicator, list_indicators
    from .security import (
        ROLE_ADMIN,
        ROLE_CLINICAL,
        ROLE_SUPERVISOR,
        TOKEN_TYPE,
        authenticate_user,
        create_access_token,
        current_user,
        require_roles,
        user_response,
    )
except ImportError:
    from schemas import (
        DataActivateRequest,
        DataCurrentResponse,
        DataUploadPreviewResponse,
        FilterOptionsResponse,
        HealthResponse,
        LoginRequest,
        LoginResponse,
        AuditEventsResponse,
        OmisosResponse,
        ReportSummary,
        SearchDNIResult,
        UploadHistoryResponse,
    )
    from services import build_report_summary, get_filter_options, load_sample_data, search_by_dni, validate_data_file
    from indicators.registry import get_indicator, list_indicators
    from security import (
        ROLE_ADMIN,
        ROLE_CLINICAL,
        ROLE_SUPERVISOR,
        TOKEN_TYPE,
        authenticate_user,
        create_access_token,
        current_user,
        require_roles,
        user_response,
    )

app = FastAPI(
    title="Sistema de Seguimiento Neonatal - MC-03",
    description="API para monitoreo de vacunacion, CRED, tamizaje neonatal y carga de datos.",
    version="0.1.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

DATA_FILE = Path(__file__).resolve().parent.parent / "data_samples" / "MC 03_FT_BCG_HVB_PAQUETE RN.xlsx"
DATA_FILES = {
    "mc03": DATA_FILE,
    "mc02": Path(__file__).resolve().parent.parent / "data_samples" / "MC 02_FT MC_02 _INFANTIL.xlsx",
}
UPLOADS_DIR = Path(__file__).resolve().parent / "uploads"
PROCESSED_UPLOADS_DIR = Path(__file__).resolve().parent / "processed_uploads"
DATA_STATE_FILE = Path(__file__).resolve().parent / "data_state.json"
JOB_STATUS_QUEUED = "queued"
JOB_STATUS_PROCESSING = "processing"
JOB_STATUS_ACTIVATED = "activated"
JOB_STATUS_FAILED = "failed"


def _json_default(value):
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _error_message(exc: Exception) -> str:
    if isinstance(exc, HTTPException):
        return str(exc.detail)
    return str(exc)


def _safe_filename(filename: str) -> str:
    clean_name = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(filename).name).strip("._")
    return clean_name or "archivo_indicador.xlsx"


def _file_sha256(filepath: Path) -> str:
    digest = hashlib.sha256()
    with filepath.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _combined_file_sha256(filepaths: list[Path]) -> str:
    digest = hashlib.sha256()
    for filepath in sorted(filepaths, key=lambda item: item.name):
        digest.update(filepath.name.encode("utf-8"))
        with filepath.open("rb") as file:
            for chunk in iter(lambda: file.read(1024 * 1024), b""):
                digest.update(chunk)
    return digest.hexdigest()


def _actor_metadata(user) -> dict:
    if user is None:
        return {}
    username = getattr(user, "username", None)
    role = getattr(user, "role", None)
    display_name = getattr(user, "display_name", None)
    if username is None and isinstance(user, dict):
        username = user.get("username")
        role = user.get("role")
        display_name = user.get("display_name")
    return {
        "username": username,
        "role": role,
        "display_name": display_name,
    }


def _load_data_state() -> dict | None:
    if not DATA_STATE_FILE.exists():
        return None
    try:
        return json.loads(DATA_STATE_FILE.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return None


def _write_data_state(state: dict) -> None:
    DATA_STATE_FILE.write_text(json.dumps(state, ensure_ascii=False, indent=2, default=_json_default), encoding="utf-8")


def _indicator_definition(indicator: str):
    try:
        return get_indicator(indicator.lower())
    except KeyError:
        raise HTTPException(status_code=404, detail=f"Indicador no disponible: {indicator}")


def _indicator_store(indicator: str) -> dict:
    store = getattr(app.state, "indicator_data", {})
    return store.get(indicator.lower(), {})


def _activation_jobs() -> dict:
    if not hasattr(app.state, "activation_jobs"):
        app.state.activation_jobs = {}
    return app.state.activation_jobs


def _active_dataframe(indicator: str):
    return _indicator_store(indicator).get("data")


def _active_cutoff_date(indicator: str):
    return _indicator_store(indicator).get("cutoff_date")


def _data_current_response(indicator: str, job: dict | None = None) -> DataCurrentResponse:
    store = _indicator_store(indicator)
    active_file = store.get("active_file")
    job = job or {}
    return DataCurrentResponse(
        has_data=store.get("data") is not None,
        active_file=str(active_file) if active_file else None,
        summary=store.get("current_data_summary"),
        processing=job.get("status") in {JOB_STATUS_QUEUED, JOB_STATUS_PROCESSING},
        job_id=job.get("id"),
        job_status=job.get("status"),
        message=job.get("message"),
        error=job.get("error"),
    )


def _build_current_summary(
    filepath: Path,
    metadata: dict | None = None,
    indicator: str = "mc03",
    validation_summary: dict | None = None,
) -> dict:
    definition = _indicator_definition(indicator)
    summary = validation_summary or definition.validate_data_file(filepath)["summary"]
    return {
        "indicator_code": summary.get("indicator_code", definition.code),
        "indicator_name": summary.get("indicator_name", definition.name),
        "filename": filepath.name,
        "original_name": (metadata or {}).get("original_name") or filepath.name,
        "file_size_bytes": summary.get("file_size_bytes"),
        "file_hash": (metadata or {}).get("file_hash") or summary.get("file_hash"),
        "uploaded_at": (metadata or {}).get("uploaded_at"),
        "uploaded_by": (metadata or {}).get("uploaded_by") or summary.get("uploaded_by"),
        "activated_at": (metadata or {}).get("activated_at"),
        "activated_by": (metadata or {}).get("activated_by") or summary.get("activated_by"),
        "sheet_name": summary.get("sheet_name"),
        "cutoff_cell": summary.get("cutoff_cell"),
        "header_row": summary.get("header_row"),
        "cutoff_date": summary.get("cutoff_date"),
        "total_rows": summary.get("total_rows"),
        "total_columns": summary.get("total_columns"),
        "loaded_columns": summary.get("loaded_columns"),
        "required_columns_count": summary.get("required_columns_count"),
        "missing_columns": summary.get("missing_columns", []),
        "omitted_columns": summary.get("omitted_columns", []),
        "provinces": summary.get("provinces", []),
        "months": summary.get("months", []),
        "validation_label": summary.get("validation_label"),
        "status_counts": summary.get("status_counts", {}),
        "denominator_counts": summary.get("denominator_counts", {}),
        "component_counts": summary.get("component_counts", {}),
        "obs_eval_counts": summary.get("obs_eval_counts", {}),
        "insurance_counts": summary.get("insurance_counts", {}),
        "storage_format": summary.get("storage_format"),
        "source_preserved": summary.get("source_preserved"),
        "files_received": summary.get("files_received"),
        "expected_files": summary.get("expected_files"),
        "subindicators_found": summary.get("subindicators_found", []),
        "missing_subindicators": summary.get("missing_subindicators", []),
        "subindicators": summary.get("subindicators", {}),
        "total_loaded_columns": summary.get("total_loaded_columns"),
    }


def _activate_data_file(
    filepath: Path,
    metadata: dict | None = None,
    indicator: str = "mc03",
    prepared_bundle: dict | None = None,
    validation_summary: dict | None = None,
) -> None:
    definition = _indicator_definition(indicator)
    sample_bundle = prepared_bundle or definition.load_sample_data(filepath)
    if sample_bundle is None:
        raise HTTPException(status_code=404, detail="No se encontro el archivo de datos")

    app.state.indicator_data[indicator] = {
        "data": sample_bundle["data"],
        "cutoff_date": sample_bundle["cutoff_date"],
        "active_file": filepath,
        "current_data_summary": _build_current_summary(filepath, metadata, indicator, validation_summary),
    }

    if indicator == "mc03":
        app.state.data = sample_bundle["data"]
        app.state.cutoff_date = sample_bundle["cutoff_date"]
        app.state.active_file = filepath
        app.state.current_data_summary = app.state.indicator_data[indicator]["current_data_summary"]


def _filter_omisos(omisos: list[dict], month: str | None = None, subindicator: str | None = None) -> list[dict]:
    filtered = omisos
    if month:
        filtered = [item for item in filtered if item.get("Mes_eva") == month]
    if subindicator:
        filtered = [item for item in filtered if item.get("subindicator_code") == subindicator]
    return filtered


def _indicator_preparer(definition):
    return getattr(definition.module, "prepare_data_file", None)


def _indicator_package_preparer(definition):
    return getattr(definition.module, "prepare_data_files", None)


def _indicator_persister(definition):
    return getattr(definition.module, "persist_active_upload", None)


def _indicator_db_searcher(definition):
    return getattr(definition.module, "search_active_by_dni", None)


def _indicator_db_reporter(definition):
    return getattr(definition.module, "build_active_report_summary", None)


def _indicator_active_upload_checker(definition):
    return getattr(definition.module, "active_upload_id", None)


def _persist_active_upload(definition, data_bundle: dict, filepath: Path, metadata: dict, validation_summary: dict | None):
    persister = _indicator_persister(definition)
    if persister is None:
        return None
    try:
        from .db.session import SessionLocal
    except ImportError:
        from db.session import SessionLocal

    with SessionLocal() as db:
        return persister(
            db,
            data_bundle["data"],
            data_bundle.get("cutoff_date"),
            filepath,
            metadata,
            validation_summary,
        )


def _search_active_upload_by_dni(definition, dni: str, province: str | None):
    searcher = _indicator_db_searcher(definition)
    active_checker = _indicator_active_upload_checker(definition)
    if searcher is None or active_checker is None:
        return None, False
    try:
        from .db.session import SessionLocal
    except ImportError:
        from db.session import SessionLocal
    try:
        from sqlalchemy.exc import SQLAlchemyError
    except ImportError:
        SQLAlchemyError = Exception

    try:
        with SessionLocal() as db:
            if active_checker(db) is None:
                return None, False
            return searcher(db, dni, province), True
    except SQLAlchemyError:
        return None, False


def _build_active_report_summary(definition, province: str | None, target: float | None = None):
    reporter = _indicator_db_reporter(definition)
    active_checker = _indicator_active_upload_checker(definition)
    if reporter is None or active_checker is None:
        return None, False
    try:
        from .db.session import SessionLocal
    except ImportError:
        from db.session import SessionLocal
    try:
        from sqlalchemy.exc import SQLAlchemyError
    except ImportError:
        SQLAlchemyError = Exception

    try:
        with SessionLocal() as db:
            if active_checker(db) is None:
                return None, False
            return reporter(db, province, target), True
    except SQLAlchemyError:
        return None, False


def _report_summary_for_request(indicator: str, province: str | None, target: float | None = None) -> tuple[str, dict]:
    definition = _indicator_definition(indicator)
    indicator = definition.code
    db_summary, used_database = _build_active_report_summary(definition, province, target)
    if used_database and db_summary is not None:
        return indicator, db_summary

    data = _active_dataframe(indicator)
    if data is None:
        raise HTTPException(status_code=404, detail="No hay datos cargados en el servidor")
    return indicator, definition.build_report_summary(data, _active_cutoff_date(indicator), province, target)


def _write_processed_bundle(path: Path, bundle: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    pd.to_pickle(bundle, path)


def _read_processed_bundle(path: Path) -> dict:
    return pd.read_pickle(path)


def _load_activation_payload(definition, pending: dict) -> tuple[Path, dict | None, dict | None]:
    processed_path = Path(pending["processed_path"]) if pending.get("processed_path") else None
    prepared_bundle = None
    validation_summary = None
    if processed_path:
        if not processed_path.exists():
            raise HTTPException(status_code=404, detail="No se encontro la carga procesada")
        processed_bundle = _read_processed_bundle(processed_path)
        prepared_bundle = {"data": processed_bundle["data"], "cutoff_date": processed_bundle["cutoff_date"]}
        validation_summary = processed_bundle.get("summary")
        return processed_path, prepared_bundle, validation_summary

    filepath = Path(pending["path"])
    validation = definition.validate_data_file(filepath)
    if not validation["valid"]:
        raise HTTPException(status_code=400, detail="El archivo ya no supera la validacion")
    return filepath, None, validation["summary"]


def _activate_pending_upload(
    upload_id: str,
    indicator: str,
    job_id: str | None = None,
    actor: dict | None = None,
) -> DataCurrentResponse:
    definition = _indicator_definition(indicator)
    indicator = definition.code
    job = _activation_jobs().get(job_id) if job_id else None
    actor = actor or (job or {}).get("actor") or {}
    if job:
        job.update(status=JOB_STATUS_PROCESSING, message="Procesando y activando datos", error=None)

    try:
        pending = app.state.pending_uploads.get(upload_id)
        if pending is None:
            raise HTTPException(status_code=404, detail="No se encontro la carga pendiente")
        if pending.get("indicator") != indicator:
            raise HTTPException(status_code=400, detail="La carga pendiente corresponde a otro indicador")
        if not pending["valid"]:
            raise HTTPException(status_code=400, detail="El archivo tiene errores de validacion y no puede activarse")

        filepath, prepared_bundle, validation_summary = _load_activation_payload(definition, pending)
        metadata = {
            "active_file": str(filepath),
            "original_name": pending["original_name"],
            "uploaded_at": pending["uploaded_at"],
            "uploaded_by": pending.get("uploaded_by"),
            "activated_at": datetime.now().isoformat(timespec="seconds"),
            "activated_by": actor.get("username") or pending.get("uploaded_by"),
            "actor_role": actor.get("role") or pending.get("uploaded_role"),
            "file_hash": pending.get("file_hash"),
        }
        db_upload_id = None
        if _indicator_persister(definition):
            active_bundle = prepared_bundle or definition.load_sample_data(filepath)
            db_upload_id = _persist_active_upload(definition, active_bundle, filepath, metadata, validation_summary)
            prepared_bundle = active_bundle
        _activate_data_file(filepath, metadata, indicator, prepared_bundle, validation_summary)
        if db_upload_id:
            app.state.indicator_data[indicator]["db_upload_id"] = str(db_upload_id)
        if indicator == "mc03":
            _write_data_state(metadata)
        app.state.pending_uploads.pop(upload_id, None)
        if job:
            job.update(
                status=JOB_STATUS_ACTIVATED,
                message="Archivo activado correctamente",
                finished_at=datetime.now().isoformat(timespec="seconds"),
                error=None,
            )
        return _data_current_response(indicator, job)
    except Exception as exc:
        error = _error_message(exc)
        if job:
            job.update(
                status=JOB_STATUS_FAILED,
                message="No se pudo activar la carga",
                error=error,
                finished_at=datetime.now().isoformat(timespec="seconds"),
            )
        raise


def _run_activation_job(job_id: str) -> None:
    job = _activation_jobs().get(job_id)
    if not job:
        return
    try:
        _activate_pending_upload(job["upload_id"], job["indicator"], job_id, job.get("actor"))
    except Exception:
        pass


def _delete_file_quietly(path: Path | None) -> None:
    if path is None:
        return
    try:
        if path.exists():
            path.unlink()
    except OSError:
        pass


def _incumplidos_xlsx(omisos: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Incumplidos"

    columns = [
        ("Subindicador", "subindicator_code"),
        ("Nombre subindicador", "subindicator_name"),
        ("Mes evaluacion", "Mes_eva"),
        ("Mes", "month"),
        ("Anio", "year"),
        ("DNI", "afi_DNI"),
        ("CNV", "NumCNV"),
        ("Fecha de nacimiento", "fec_Nac"),
        ("Nombres", "afi_nombres"),
        ("Apellido paterno", "afi_appaterno"),
        ("Apellido materno", "afi_apmaterno"),
        ("Provincia", "Desc_prov"),
        ("Microred", "Des_MicroRed"),
        ("Codigo RENAES", "pre_CodigoRENAES"),
        ("Establecimiento", "Des_EESS"),
        ("Componentes observados", "components_observed"),
        ("Fecha atencion observada", "attention_date"),
        ("Edad atencion dias", "attention_age_days"),
        ("EESS atencion", "attention_facility"),
        ("Profesional atencion", "attention_professional"),
        ("Alertas clinicas", "clinical_alerts"),
        ("Motivo de incumplimiento", "reason"),
    ]

    for column_index, (header, _) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F172A")

    for row_index, item in enumerate(omisos, start=2):
        for column_index, (_, key) in enumerate(columns, start=1):
            value = item.get(key)
            if isinstance(value, list):
                value = "; ".join(str(part) for part in value if part)
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            if key in {"components_observed", "reason", "clinical_alerts"}:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_index, (header, key) in enumerate(columns, start=1):
        max_length = len(header)
        for item in omisos:
            max_length = max(max_length, len(str(item.get(key) or "")))
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max_length + 2, 55)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@app.on_event("startup")
def startup_event():
    UPLOADS_DIR.mkdir(parents=True, exist_ok=True)
    PROCESSED_UPLOADS_DIR.mkdir(parents=True, exist_ok=True)
    app.state.pending_uploads = {}
    app.state.activation_jobs = {}
    app.state.indicator_data = {}

    state = _load_data_state()
    active_file = Path(state["active_file"]) if state and state.get("active_file") else DATA_FILE
    if not active_file.exists():
        active_file = DATA_FILE

    try:
        _activate_data_file(active_file, state, "mc03")
    except Exception:
        app.state.data = None
        app.state.cutoff_date = None
        app.state.active_file = None
        app.state.current_data_summary = None

    for indicator, filepath in DATA_FILES.items():
        if indicator == "mc03" or not filepath.exists():
            continue
        try:
            _activate_data_file(filepath, None, indicator)
        except Exception:
            app.state.indicator_data[indicator] = {
                "data": None,
                "cutoff_date": None,
                "active_file": None,
                "current_data_summary": None,
            }


@app.get("/health", response_model=HealthResponse, tags=["Sistema"])
def health_check():
    source_file = str(app.state.active_file) if getattr(app.state, "active_file", None) else str(DATA_FILE)
    return HealthResponse(status="ok", version="0.1.0", source_file=source_file)


@app.post("/api/auth/login", response_model=LoginResponse, tags=["Seguridad"])
def api_auth_login(payload: LoginRequest):
    user = authenticate_user(payload.username, payload.password)
    if user is None:
        raise HTTPException(status_code=401, detail="Usuario o contrasena incorrectos")
    return LoginResponse(access_token=create_access_token(user), token_type=TOKEN_TYPE, user=user_response(user))


@app.get("/api/auth/me", tags=["Seguridad"])
def api_auth_me(user=Depends(current_user)):
    return user_response(user)


@app.get("/api/indicators", tags=["Sistema"])
def api_indicators(_user=Depends(require_roles(ROLE_CLINICAL, ROLE_SUPERVISOR, ROLE_ADMIN))):
    return [
        {
            "code": indicator.code,
            "name": indicator.name,
            "has_data": _indicator_store(indicator.code).get("data") is not None,
        }
        for indicator in list_indicators()
    ]


@app.get("/api/data/current", response_model=DataCurrentResponse, tags=["Carga de datos"])
def api_data_current(indicator: str = Query("mc03"), _user=Depends(require_roles(ROLE_ADMIN))):
    definition = _indicator_definition(indicator)
    indicator = definition.code
    return _data_current_response(indicator)


@app.get("/api/data/uploads", response_model=UploadHistoryResponse, tags=["Carga de datos"])
def api_data_upload_history(
    indicator: str = Query("mc03"),
    limit: int = Query(30, ge=1, le=100),
    _user=Depends(require_roles(ROLE_ADMIN)),
):
    definition = _indicator_definition(indicator)
    try:
        from .db.session import SessionLocal
        from .db.audit import list_upload_history
    except ImportError:
        from db.session import SessionLocal
        from db.audit import list_upload_history

    with SessionLocal() as db:
        return UploadHistoryResponse(uploads=list_upload_history(db, definition.code, limit))


@app.get("/api/audit/events", response_model=AuditEventsResponse, tags=["Seguridad"])
def api_audit_events(
    indicator: str | None = Query(None),
    upload_id: str | None = Query(None),
    limit: int = Query(100, ge=1, le=500),
    _user=Depends(require_roles(ROLE_ADMIN)),
):
    indicator_code = _indicator_definition(indicator).code if isinstance(indicator, str) and indicator else None
    parsed_upload_id = UUID(upload_id) if isinstance(upload_id, str) and upload_id else None
    try:
        from .db.session import SessionLocal
        from .db.audit import list_audit_events
    except ImportError:
        from db.session import SessionLocal
        from db.audit import list_audit_events

    with SessionLocal() as db:
        return AuditEventsResponse(events=list_audit_events(db, indicator_code, parsed_upload_id, limit))


@app.post("/api/data/upload-preview", response_model=DataUploadPreviewResponse, tags=["Carga de datos"])
def api_data_upload_preview(
    file: UploadFile | None = File(None),
    files: list[UploadFile] = File(default=[]),
    indicator: str = Query("mc03"),
    _user=Depends(require_roles(ROLE_ADMIN)),
):
    definition = _indicator_definition(indicator)
    indicator = definition.code
    uploaded_files = [item for item in ([file] if file else []) + (files or []) if item and item.filename]
    if not uploaded_files:
        raise HTTPException(status_code=400, detail="Debe cargar al menos un archivo Excel .xlsx")
    if any(Path(item.filename or "").suffix.lower() != ".xlsx" for item in uploaded_files):
        raise HTTPException(status_code=400, detail="Solo se permiten archivos Excel .xlsx")

    upload_id = uuid4().hex
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    destinations: list[Path] = []
    for index, uploaded_file in enumerate(uploaded_files, start=1):
        original_filename = _safe_filename(uploaded_file.filename or f"archivo_{definition.code}_{index}.xlsx")
        destination = UPLOADS_DIR / f"{timestamp}_{upload_id[:8]}_{index}_{original_filename}"
        with destination.open("wb") as buffer:
            shutil.copyfileobj(uploaded_file.file, buffer)
        destinations.append(destination)

    package_preparer = _indicator_package_preparer(definition)
    package_upload = package_preparer is not None
    if len(destinations) > 1 and package_preparer is None:
        for destination in destinations:
            _delete_file_quietly(destination)
        raise HTTPException(status_code=400, detail="El indicador seleccionado no acepta carga multiparchivo")

    destination = destinations[0]
    file_hash = _combined_file_sha256(destinations) if len(destinations) > 1 else _file_sha256(destination)
    actor = _actor_metadata(_user)
    original_name = (
        f"{definition.code}_paquete_{len(destinations)}_archivos.xlsx"
        if package_upload
        else (uploaded_files[0].filename or f"archivo_{definition.code}.xlsx")
    )

    processed_path = None
    preparer = package_preparer or _indicator_preparer(definition)
    if preparer:
        prepared = package_preparer(destinations) if package_preparer else preparer(destination)
        validation = prepared["validation"]
        if validation["valid"]:
            validation["summary"]["storage_format"] = "processed_package" if package_upload else "processed_dataframe"
            validation["summary"]["source_preserved"] = False
            validation["summary"]["file_size_bytes"] = sum(path.stat().st_size for path in destinations if path.exists())
            processed_path = PROCESSED_UPLOADS_DIR / f"{timestamp}_{upload_id[:8]}_{indicator}.pkl"
            _write_processed_bundle(
                processed_path,
                {
                    "data": prepared["data"],
                    "cutoff_date": prepared.get("cutoff_date") or prepared.get("cutoff_dates"),
                    "summary": validation["summary"],
                },
            )
        for destination_item in destinations:
            _delete_file_quietly(destination_item)
    else:
        if package_upload:
            for destination_item in destinations:
                _delete_file_quietly(destination_item)
            raise HTTPException(status_code=400, detail="El indicador seleccionado no tiene preparador multiparchivo")
        validation = definition.validate_data_file(destination)

    uploaded_at = datetime.now().isoformat(timespec="seconds")
    summary = {
        **validation["summary"],
        "filename": destination.name,
        "original_name": original_name,
        "uploaded_at": uploaded_at,
        "uploaded_by": actor.get("username"),
        "file_hash": file_hash,
    }

    app.state.pending_uploads[upload_id] = {
        "path": str(destination),
        "paths": [str(path) for path in destinations],
        "processed_path": str(processed_path) if processed_path else None,
        "original_name": original_name,
        "uploaded_at": uploaded_at,
        "uploaded_by": actor.get("username"),
        "uploaded_role": actor.get("role"),
        "file_hash": file_hash,
        "valid": validation["valid"],
        "indicator": indicator,
    }

    return DataUploadPreviewResponse(
        upload_id=upload_id,
        valid=validation["valid"],
        errors=validation["errors"],
        warnings=validation["warnings"],
        summary=summary,
    )


@app.post("/api/data/activate", response_model=DataCurrentResponse, tags=["Carga de datos"])
def api_data_activate(
    payload: DataActivateRequest,
    background_tasks: BackgroundTasks,
    indicator: str = Query("mc03"),
    _user=Depends(require_roles(ROLE_ADMIN)),
):
    definition = _indicator_definition(indicator)
    indicator = definition.code
    pending = app.state.pending_uploads.get(payload.upload_id)
    if pending is None:
        raise HTTPException(status_code=404, detail="No se encontro la carga pendiente")
    if pending.get("indicator") != indicator:
        raise HTTPException(status_code=400, detail="La carga pendiente corresponde a otro indicador")
    if not pending["valid"]:
        raise HTTPException(status_code=400, detail="El archivo tiene errores de validacion y no puede activarse")

    if _indicator_persister(definition):
        existing_job_id = pending.get("activation_job_id")
        jobs = _activation_jobs()
        existing_job = jobs.get(existing_job_id) if existing_job_id else None
        if existing_job and existing_job.get("status") in {JOB_STATUS_QUEUED, JOB_STATUS_PROCESSING}:
            return _data_current_response(indicator, existing_job)

        job_id = uuid4().hex
        job = {
            "id": job_id,
            "indicator": indicator,
            "upload_id": payload.upload_id,
            "actor": _actor_metadata(_user),
            "status": JOB_STATUS_QUEUED,
            "message": "Activacion en cola",
            "error": None,
            "started_at": datetime.now().isoformat(timespec="seconds"),
            "finished_at": None,
        }
        jobs[job_id] = job
        pending["activation_job_id"] = job_id
        background_tasks.add_task(_run_activation_job, job_id)
        return _data_current_response(indicator, job)

    return _activate_pending_upload(payload.upload_id, indicator, actor=_actor_metadata(_user))

@app.get("/api/data/activation/{job_id}", response_model=DataCurrentResponse, tags=["Carga de datos"])
def api_data_activation_status(
    job_id: str,
    indicator: str = Query("mc03"),
    _user=Depends(require_roles(ROLE_ADMIN)),
):
    definition = _indicator_definition(indicator)
    indicator = definition.code
    job = _activation_jobs().get(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="No se encontro el trabajo de activacion")
    if job.get("indicator") != indicator:
        raise HTTPException(status_code=400, detail="El trabajo de activacion corresponde a otro indicador")
    return _data_current_response(indicator, job)


@app.get("/api/config/options", response_model=FilterOptionsResponse, tags=["Configuracion"])
def api_config_options(indicator: str = Query("mc03"), _user=Depends(require_roles(ROLE_ADMIN))):
    data = _active_dataframe(indicator)
    if data is None:
        raise HTTPException(status_code=404, detail="No hay datos cargados en el servidor")

    definition = _indicator_definition(indicator)
    return FilterOptionsResponse(**definition.get_filter_options(data))


@app.get("/api/report/summary", response_model=ReportSummary, tags=["Reporte"])
def api_report_summary(
    province: str = Query("ABANCAY"),
    target: float = Query(70.7),
    indicator: str = Query("mc03"),
    _user=Depends(require_roles(ROLE_SUPERVISOR, ROLE_ADMIN)),
):
    _, summary_dict = _report_summary_for_request(indicator, province, target)
    return ReportSummary(**summary_dict)


@app.get("/api/report/omisos", response_model=OmisosResponse, tags=["Reporte"])
def api_report_omisos(
    province: str = Query("ABANCAY"),
    indicator: str = Query("mc03"),
    subindicator: str | None = Query(None),
    _user=Depends(require_roles(ROLE_SUPERVISOR, ROLE_ADMIN)),
):
    _, summary = _report_summary_for_request(indicator, province)
    return OmisosResponse(omisos=_filter_omisos(summary["omisos"], subindicator=subindicator))


@app.get("/api/report/incumplidos", response_model=OmisosResponse, tags=["Reporte"])
def api_report_incumplidos(
    province: str = Query("ABANCAY"),
    indicator: str = Query("mc03"),
    subindicator: str | None = Query(None),
    _user=Depends(require_roles(ROLE_SUPERVISOR, ROLE_ADMIN)),
):
    _, summary = _report_summary_for_request(indicator, province)
    return OmisosResponse(omisos=_filter_omisos(summary["omisos"], subindicator=subindicator))


@app.get("/api/report/omisos.csv", tags=["Reporte"])
def api_report_omisos_csv(
    province: str = Query("ABANCAY"),
    indicator: str = Query("mc03"),
    subindicator: str | None = Query(None),
    _user=Depends(require_roles(ROLE_SUPERVISOR, ROLE_ADMIN)),
):
    indicator, summary = _report_summary_for_request(indicator, province)
    output = StringIO()
    pd.DataFrame(_filter_omisos(summary["omisos"], subindicator=subindicator)).to_csv(output, index=False)
    output.seek(0)
    filename_subindicator = f"_{subindicator}" if subindicator else ""
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="omisos_{indicator}{filename_subindicator}.csv"'},
    )


@app.get("/api/report/incumplidos.csv", tags=["Reporte"])
def api_report_incumplidos_csv(
    province: str = Query("ABANCAY"),
    month: str | None = Query(None),
    subindicator: str | None = Query(None),
    indicator: str = Query("mc03"),
    _user=Depends(require_roles(ROLE_SUPERVISOR, ROLE_ADMIN)),
):
    indicator, summary = _report_summary_for_request(indicator, province)
    output = StringIO()
    pd.DataFrame(_filter_omisos(summary["omisos"], month, subindicator)).to_csv(output, index=False)
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="incumplidos_{indicator}{f"_{subindicator}" if subindicator else ""}.csv"'},
    )


@app.get("/api/report/incumplidos.xlsx", tags=["Reporte"])
def api_report_incumplidos_xlsx(
    province: str = Query("ABANCAY"),
    month: str | None = Query(None),
    subindicator: str | None = Query(None),
    indicator: str = Query("mc03"),
    _user=Depends(require_roles(ROLE_SUPERVISOR, ROLE_ADMIN)),
):
    indicator, summary = _report_summary_for_request(indicator, province)
    content = _incumplidos_xlsx(_filter_omisos(summary["omisos"], month, subindicator))
    filename_month = f"_{month}" if month else ""
    filename_subindicator = f"_{subindicator}" if subindicator else ""
    return StreamingResponse(
        iter([content]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="incumplidos_{indicator}{filename_subindicator}{filename_month}.xlsx"'},
    )


@app.get("/api/search/dni/{dni}", response_model=SearchDNIResult, tags=["Busqueda"])
def api_search_dni(
    dni: str,
    province: str = Query("ABANCAY"),
    indicator: str = Query("mc03"),
    _user=Depends(require_roles(ROLE_CLINICAL, ROLE_SUPERVISOR, ROLE_ADMIN)),
):
    definition = _indicator_definition(indicator)
    indicator = definition.code
    db_result, searched_database = _search_active_upload_by_dni(definition, dni, province)
    if searched_database:
        if db_result is None:
            raise HTTPException(status_code=404, detail="DNI no encontrado en la carga activa")
        return SearchDNIResult(**db_result)

    data = _active_dataframe(indicator)
    if data is None:
        raise HTTPException(status_code=404, detail="No hay datos cargados en el servidor")

    result = definition.search_by_dni(data, dni, _active_cutoff_date(indicator), province)
    if result is None:
        raise HTTPException(status_code=404, detail=f"No se encontro registro para DNI: {dni}")

    return SearchDNIResult(**result)
