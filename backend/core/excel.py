from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from .dates import parse_excel_date


def open_workbook(filepath: Path):
    return load_workbook(filepath, data_only=True, read_only=True)


def read_sheet(filepath: Path, sheet_name: str, header_row: int, usecols: list[str] | set[str] | None = None) -> pd.DataFrame:
    selected_columns = {str(column).strip() for column in usecols or [] if column}
    usecols_arg = (lambda column: str(column).strip() in selected_columns) if selected_columns else None
    df = pd.read_excel(
        filepath,
        sheet_name=sheet_name,
        header=header_row - 1,
        engine="openpyxl",
        usecols=usecols_arg,
    )
    df.columns = [str(column).strip() for column in df.columns]
    return df


def read_header_values(filepath: Path, sheet_name: str, header_row: int) -> list[str]:
    workbook = open_workbook(filepath)
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"La hoja '{sheet_name}' no se encuentra en el archivo de datos")
    sheet = workbook[sheet_name]
    values = next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
    workbook.close()
    return [str(value).strip() for value in values if value is not None and str(value).strip()]


def read_cutoff_date(filepath: Path, sheet_name: str, cell: str) -> Any:
    workbook = open_workbook(filepath)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"La hoja '{sheet_name}' no se encuentra en el archivo de datos")
    return parse_excel_date(workbook[sheet_name][cell].value)


def build_validation_error(message: str, summary: dict[str, Any] | None = None) -> dict[str, Any]:
    return {"valid": False, "errors": [message], "warnings": [], "summary": summary or {}}
